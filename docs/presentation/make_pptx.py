# dev-briefing.md -> dev-briefing.pptx 생성기.
#
# md 계약 요약 (정본: docs/work/20260809-dev-briefing/req-design.md DES-01, v2):
#   - 문서 첫 '# ' 줄: 덱 제목(표지). 첫 '##' 전의 '- ' 줄은 표지 부제, '> 노트:'는 표지 노트.
#   - 각 '## ' 줄: 본문 슬라이드 1장.
#   - '- ' / '  - ' : 불릿(2단까지). '> 노트:'(연속 '>' 줄 포함): 발표자 노트.
#   - ``` 펜스: 태그로 렌더 유형 지정 — ```flow(세로 흐름도) ```rows(체브런 행)
#     ```ladder(번호 스텝), 무태그는 고정폭 패널. 다른 태그·요소(표·### 등)는 오류로 중단.
#   - 디자인(DES-05 팔레트·배치)과 카드 배치(DES-06)는 스크립트가 자동 적용한다.
#
# 실행: python make_pptx.py [입력.md] [출력.pptx]  (인자 없으면 같은 폴더의 dev-briefing.*)
# 생성 직후 결과 파일을 다시 열어 슬라이드 수·제목·노트를 md와 대조한다(AC-02).
import math
import os
import sys

KOREAN_FONT = "맑은 고딕"
CODE_FONT = "Consolas"

# DES-05 디자인 시스템 — 저채도·고대비 5계열 팔레트 (정본: req-design.md)
INK = (0x24, 0x30, 0x4A)  # 제목·본문 텍스트, 표지·마무리 배경
BG = (0xFA, 0xF9, 0xF5)  # 본문 슬라이드 배경
ACCENT = (0x1B, 0x7A, 0x6E)  # 틸 — 제목 언더바, 다이어그램 기본
ACCENT_LIGHT = (0xE3, 0xF0, 0xED)
GATE = (0xC0, 0x5B, 0x2E)  # 번트 오렌지 — 승인 관문(★)·경고성 강조
GATE_LIGHT = (0xF7, 0xE7, 0xDC)
GRAY = (0x6A, 0x74, 0x86)
PANEL = (0xF0, 0xEE, 0xE7)  # 코드·트리 패널
WHITE = (0xFF, 0xFF, 0xFF)
COVER_SUB = (0xD9, 0xDE, 0xE8)

FENCE_TAGS = ("flow", "rows", "ladder")

# DES-06 카드형 배치 — 제목 → 템플릿 매핑. 구조가 안 맞으면 불릿으로 폴백(fail-soft).
CARD_TEMPLATES = {
    "AI 코딩 위임의 실패 모드": "failure-cards",
    "세션 핸드오프 — 개념과 3층 설계": "layer-bands",
}

MARGIN = 0.6  # DES-05 일관 마진(인치)
CONTENT_W = 13.333 - 2 * MARGIN
BODY_TOP = 1.55


class MdContractError(Exception):
    pass


class BodySlide:
    def __init__(self, title):
        self.title = title
        # ("bullet", level, text) | ("code"|"flow"|"rows"|"ladder", text)
        self.blocks = []
        self.notes = ""


class Deck:
    def __init__(self):
        self.title = None
        self.cover_lines = []
        self.cover_notes = ""
        self.slides = []


def parse_md(text):
    deck = Deck()
    current = None  # None = 표지 영역, 이후 BodySlide
    code_lines = None
    fence_tag = None
    in_notes = False

    for lineno, raw in enumerate(text.splitlines(), 1):
        line = raw.rstrip()

        if code_lines is not None:
            if line.startswith("```"):
                current.blocks.append((fence_tag, "\n".join(code_lines)))
                code_lines = None
                fence_tag = None
            else:
                code_lines.append(raw)
            continue

        if not line.strip():
            in_notes = False
            continue

        if line.startswith("## "):
            current = BodySlide(line[3:].strip())
            deck.slides.append(current)
            in_notes = False
        elif line.startswith("# ") and deck.title is None and current is None:
            deck.title = line[2:].strip()
        elif line.startswith("```"):
            if current is None:
                raise MdContractError("line %d: code block not allowed on cover" % lineno)
            tag = line[3:].strip()
            if tag and tag not in FENCE_TAGS:
                raise MdContractError(
                    "line %d: unsupported fence tag: %r" % (lineno, tag)
                )
            fence_tag = tag or "code"
            code_lines = []
        elif line.startswith("> 노트:"):
            note = line[len("> 노트:"):].strip()
            if current is None:
                deck.cover_notes = (deck.cover_notes + "\n" + note).strip()
            else:
                current.notes = (current.notes + "\n" + note).strip()
            in_notes = True
        elif line.startswith(">") and in_notes:
            note = line.lstrip(">").strip()
            if current is None:
                deck.cover_notes += "\n" + note
            else:
                current.notes += "\n" + note
        elif line.startswith("- "):
            if current is None:
                deck.cover_lines.append(line[2:].strip())
            else:
                current.blocks.append(("bullet", 0, line[2:].strip()))
            in_notes = False
        elif line.startswith("  - "):
            if current is None:
                raise MdContractError("line %d: nested bullet not allowed on cover" % lineno)
            current.blocks.append(("bullet", 1, line[4:].strip()))
            in_notes = False
        else:
            raise MdContractError(
                "line %d: unsupported construct: %r" % (lineno, line[:60])
            )

    if code_lines is not None:
        raise MdContractError("unterminated code fence")
    if deck.title is None:
        raise MdContractError("missing deck title ('# ...') on first line")
    if not deck.slides:
        raise MdContractError("no body slides ('## ...') found")
    return deck


def _style_paragraph(para, size_pt, mono=False, bullet=True, color=None, bold=False):
    from lxml import etree
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
    from pptx.util import Pt

    if not bullet:
        pPr = para._p.get_or_add_pPr()
        for tag in ("a:buChar", "a:buAutoNum", "a:buNone"):
            for el in pPr.findall(qn(tag)):
                pPr.remove(el)
        etree.SubElement(pPr, qn("a:buNone"))

    runs = para.runs or [para.add_run()]
    for run in runs:
        run.font.size = Pt(size_pt)
        run.font.name = CODE_FONT if mono else KOREAN_FONT
        if color is not None:
            run.font.color.rgb = RGBColor(*color)
        if bold:
            run.font.bold = True
        rPr = run._r.get_or_add_rPr()
        ea = rPr.find(qn("a:ea"))
        if ea is None:
            ea = etree.SubElement(rPr, qn("a:ea"))
        ea.set("typeface", KOREAN_FONT)


def _bullet_char(para, level):
    # 자유 텍스트 상자에 자리표시자와 같은 불릿 마커·들여쓰기를 적용한다.
    from lxml import etree
    from pptx.oxml.ns import qn

    pPr = para._p.get_or_add_pPr()
    pPr.set("marL", "274320" if level == 0 else "548640")
    pPr.set("indent", "-274320")
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", "•" if level == 0 else "–")


def _fill_body(text_frame, blocks, text_color=INK):
    first = True
    for block in blocks:
        if block[0] == "bullet":
            _, level, text = block
            para = text_frame.paragraphs[0] if first else text_frame.add_paragraph()
            first = False
            para.text = text
            para.level = level
            _style_paragraph(para, 20 if level == 0 else 17, color=text_color)
        else:  # code — 무태그 펜스(패널 사각형은 호출부가 뒤에 깐다)
            for code_line in block[1].split("\n"):
                para = text_frame.paragraphs[0] if first else text_frame.add_paragraph()
                first = False
                para.text = code_line
                _style_paragraph(para, 13, mono=True, bullet=False, color=INK)


def _place(shape, left, top, width, height):
    from pptx.util import Inches

    shape.left = Inches(left)
    shape.top = Inches(top)
    shape.width = Inches(width)
    shape.height = Inches(height)


def _set_bg(slide, rgb):
    from pptx.dml.color import RGBColor

    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*rgb)


def _paint(shape, fill_rgb, line_rgb=None):
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*fill_rgb)
    if line_rgb is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = RGBColor(*line_rgb)
        shape.line.width = Pt(1)
    shape.shadow.inherit = False


def _fit(size_pt, text):
    # 긴 문구의 넘침 완화 — NFR-04 폰트 축소 규칙.
    n = len(text)
    if n > 60:
        return size_pt - 4
    if n > 38:
        return size_pt - 2
    return size_pt


def _shape_text(shape, text, size_pt, color=INK, bold=False, center=True, mono=False):
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches

    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, line in enumerate(text.split("\n")):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = line
        if center:
            para.alignment = PP_ALIGN.CENTER
        _style_paragraph(para, size_pt, mono=mono, bullet=False, color=color, bold=bold)


def _send_back(slide, shape):
    spTree = slide.shapes._spTree
    el = shape._element
    spTree.remove(el)
    spTree.insert(2, el)


def _drop(placeholder):
    el = placeholder._element
    el.getparent().remove(el)


def _est_bullet_height(bullets):
    # 자동 줄바꿈을 포함한 불릿 높이 추정(인치) — 패널·순차 배치용.
    h = 0.0
    for level, text in bullets:
        per_line = 44 if level == 0 else 50
        lines = max(1, math.ceil(len(text) / per_line))
        h += lines * (0.40 if level == 0 else 0.34)
    return h + 0.06


def _panels_behind_code(slide, blocks):
    # 본문 자리표시자 텍스트 흐름 속 코드 줄 뒤에 연회색 패널을 깐다.
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    y = BODY_TOP + 0.06
    for block in blocks:
        if block[0] == "bullet":
            y += _est_bullet_height([(block[1], block[2])]) - 0.06
        else:
            n = len(block[1].split("\n"))
            h = n * 0.22 + 0.12
            rect = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(MARGIN + 0.05), Inches(y - 0.05),
                Inches(CONTENT_W - 0.1), Inches(h),
            )
            _paint(rect, PANEL)
            _send_back(slide, rect)
            y += n * 0.22


def _add_flow(slide, text, top):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    lines = [l.strip() for l in text.split("\n") if l.strip()]
    box_w, box_h, arrow_h, gap = 7.6, 0.46, 0.14, 0.05
    x = MARGIN + (CONTENT_W - box_w) / 2
    y = top
    for i, line in enumerate(lines):
        if i:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.DOWN_ARROW,
                Inches(MARGIN + CONTENT_W / 2 - 0.09), Inches(y),
                Inches(0.18), Inches(arrow_h),
            )
            _paint(arrow, GRAY)
            y += arrow_h + gap
        gate = "★" in line
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x), Inches(y), Inches(box_w), Inches(box_h),
        )
        _paint(box, GATE_LIGHT if gate else ACCENT_LIGHT, GATE if gate else ACCENT)
        _shape_text(box, line, _fit(16, line), color=INK, bold=gate)
        y += box_h + gap
    return y - top


def _add_rows(slide, text, top):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    lines = [l.strip() for l in text.split("\n") if l.strip()]
    row_h, row_gap, seg_gap = 0.55, 0.15, 0.08
    y = top
    for line in lines:
        segs = [s.strip() for s in line.split("→")]
        seg_w = (CONTENT_W - seg_gap * (len(segs) - 1)) / len(segs)
        x = MARGIN
        for j, seg in enumerate(segs):
            gate = "★" in seg
            shp = slide.shapes.add_shape(
                MSO_SHAPE.PENTAGON if j == 0 else MSO_SHAPE.CHEVRON,
                Inches(x), Inches(y), Inches(seg_w), Inches(row_h),
            )
            _paint(shp, GATE_LIGHT if gate else ACCENT_LIGHT, GATE if gate else ACCENT)
            _shape_text(shp, seg, _fit(14, seg), color=INK, bold=gate)
            x += seg_w + seg_gap
        y += row_h + row_gap
    return y - top


def _add_ladder(slide, text, top):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    lines = [l.strip() for l in text.split("\n") if l.strip()]
    step_h, gap = 0.44, 0.06
    two_col = len(lines) > 5  # 긴 사다리는 2열로 접어 세로 넘침을 막는다(NFR-04)
    rows_per_col = math.ceil(len(lines) / 2) if two_col else len(lines)
    col_w = (CONTENT_W - 0.3) / 2 if two_col else CONTENT_W
    for i, line in enumerate(lines):
        col, row = divmod(i, rows_per_col)
        x = MARGIN + col * (col_w + 0.3)
        y = top + row * (step_h + gap)
        chip = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(x + 0.05), Inches(y + 0.05),
            Inches(0.34), Inches(0.34),
        )
        _paint(chip, ACCENT)
        _shape_text(chip, str(i + 1), 13, color=WHITE, bold=True)
        txt = slide.shapes.add_textbox(
            Inches(x + 0.55), Inches(y),
            Inches(col_w - 0.55), Inches(step_h),
        )
        _shape_text(txt, line, 16, color=INK, center=False)
    return rows_per_col * (step_h + gap)


def _add_code_panel(slide, text, top):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    lines = text.split("\n")
    h = len(lines) * 0.24 + 0.2
    rect = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(MARGIN + 0.05), Inches(top),
        Inches(CONTENT_W - 0.1), Inches(h),
    )
    _paint(rect, PANEL)
    _shape_text(rect, text, 13, color=INK, center=False, mono=True)
    return h


def _add_bullet_box(slide, bullets, top, text_color):
    from pptx.util import Inches

    h = _est_bullet_height(bullets)
    box = slide.shapes.add_textbox(
        Inches(MARGIN), Inches(top), Inches(CONTENT_W), Inches(h)
    )
    tf = box.text_frame
    tf.word_wrap = True
    for i, (level, text) in enumerate(bullets):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = text
        _style_paragraph(para, 20 if level == 0 else 17, color=text_color)
        _bullet_char(para, level)
    return h


def _render_sequential(slide, blocks, text_color):
    # 다이어그램 블록이 있는 슬라이드: 블록 순서대로 위→아래 배치.
    y = BODY_TOP
    pending = []
    for block in blocks + [None]:
        if block is not None and block[0] == "bullet":
            pending.append((block[1], block[2]))
            continue
        if pending:
            y += _add_bullet_box(slide, pending, y, text_color) + 0.12
            pending = []
        if block is None:
            break
        kind, text = block
        if kind == "flow":
            y += _add_flow(slide, text, y) + 0.15
        elif kind == "rows":
            y += _add_rows(slide, text, y) + 0.15
        elif kind == "ladder":
            y += _add_ladder(slide, text, y) + 0.15
        else:  # code
            y += _add_code_panel(slide, text, y) + 0.12


def _render_cards(slide, template, blocks, text_color):
    # 카드형 배치 — 기대 구조가 아니면 False(호출부가 불릿으로 폴백).
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    if any(b[0] != "bullet" for b in blocks):
        return False
    bullets = [(b[1], b[2]) for b in blocks]

    if template == "failure-cards":
        if any(level != 0 for level, _ in bullets) or len(bullets) < 5:
            return False
        cards, conclusion = bullets[:4], bullets[4:]
        card_w = (CONTENT_W - 0.3) / 2
        card_h = 1.55
        for i, (_, text) in enumerate(cards):
            x = MARGIN + (i % 2) * (card_w + 0.3)
            y = BODY_TOP + 0.1 + (i // 2) * (card_h + 0.25)
            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(x), Inches(y), Inches(card_w), Inches(card_h),
            )
            _paint(card, WHITE, ACCENT)
            _shape_text(card, text, _fit(16, text), color=INK)
        y = BODY_TOP + 0.1 + 2 * (card_h + 0.25) + 0.1
        for _, text in conclusion:
            band = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(MARGIN), Inches(y), Inches(CONTENT_W), Inches(0.6),
            )
            _paint(band, GATE_LIGHT, GATE)
            _shape_text(band, text, _fit(16, text), color=INK, bold=True)
            y += 0.72
        return True

    if template == "layer-bands":
        prefixes = ("A층", "B층", "C층")
        band_idx = [
            i for i, (level, text) in enumerate(bullets)
            if level == 0 and text.startswith(prefixes)
        ]
        if len(band_idx) != 3:
            return False
        intro = bullets[: band_idx[0]]
        outro = bullets[band_idx[2] + 1:]
        y = BODY_TOP
        if intro:
            y += _add_bullet_box(slide, intro, y, text_color) + 0.15
        for i in band_idx:
            _, text = bullets[i]
            band = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(MARGIN), Inches(y), Inches(CONTENT_W), Inches(0.9),
            )
            _paint(band, ACCENT_LIGHT, ACCENT)
            _shape_text(band, text, _fit(16, text), color=INK, center=False)
            y += 1.02
        if outro:
            _add_bullet_box(slide, outro, y + 0.05, text_color)
        return True

    return False


def build_pptx(deck, out_path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    cover = prs.slides.add_slide(prs.slide_layouts[0])
    _set_bg(cover, INK)
    cover.shapes.title.text = deck.title
    _place(cover.shapes.title, 0.9, 2.2, 11.5, 1.6)
    for para in cover.shapes.title.text_frame.paragraphs:
        _style_paragraph(para, 40, color=WHITE, bold=True)
    subtitle = cover.placeholders[1]
    _place(subtitle, 0.9, 4.0, 11.5, 2.2)
    subtitle.text = "\n".join(deck.cover_lines)
    for para in subtitle.text_frame.paragraphs:
        _style_paragraph(para, 18, color=COVER_SUB)
    rule = cover.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.95), Inches(3.75), Inches(2.0), Inches(0.055)
    )
    _paint(rule, ACCENT)
    if deck.cover_notes:
        cover.notes_slide.notes_text_frame.text = deck.cover_notes

    for idx, spec in enumerate(deck.slides):
        closing = idx == len(deck.slides) - 1
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        _set_bg(slide, INK if closing else BG)

        slide.shapes.title.text = spec.title
        _place(slide.shapes.title, MARGIN, 0.32, CONTENT_W, 0.9)
        for para in slide.shapes.title.text_frame.paragraphs:
            _style_paragraph(para, 28, color=WHITE if closing else INK, bold=True)
        underbar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(MARGIN), Inches(1.28), Inches(2.0), Inches(0.055)
        )
        _paint(underbar, ACCENT)

        text_color = WHITE if closing else INK
        body = slide.placeholders[1]
        template = CARD_TEMPLATES.get(spec.title)
        if template and _render_cards(slide, template, spec.blocks, text_color):
            _drop(body)
        elif any(b[0] in FENCE_TAGS for b in spec.blocks):
            _drop(body)
            _render_sequential(slide, spec.blocks, text_color)
        else:
            _place(body, MARGIN, BODY_TOP, CONTENT_W, 5.6)
            _fill_body(body.text_frame, spec.blocks, text_color)
            _panels_behind_code(slide, spec.blocks)

        if spec.notes:
            slide.notes_slide.notes_text_frame.text = spec.notes

    prs.save(out_path)


def verify_pptx(deck, out_path):
    from pptx import Presentation

    problems = []
    prs = Presentation(out_path)
    slides = list(prs.slides)
    expected = 1 + len(deck.slides)
    if len(slides) != expected:
        problems.append("slide count: expected %d, got %d" % (expected, len(slides)))
        return problems

    expected_titles = [deck.title] + [s.title for s in deck.slides]
    expected_notes = [deck.cover_notes] + [s.notes for s in deck.slides]
    for i, (slide, title, notes) in enumerate(zip(slides, expected_titles, expected_notes)):
        actual = slide.shapes.title.text if slide.shapes.title else ""
        if actual != title:
            problems.append("slide %d title: expected %r, got %r" % (i + 1, title, actual))
        if notes:
            actual_notes = (
                slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
            )
            if actual_notes != notes:
                problems.append("slide %d notes mismatch" % (i + 1))
    return problems


def main(argv):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except AttributeError:
        pass

    base = os.path.dirname(os.path.abspath(__file__))
    md_path = argv[1] if len(argv) > 1 else os.path.join(base, "dev-briefing.md")
    out_path = argv[2] if len(argv) > 2 else os.path.join(base, "dev-briefing.pptx")

    try:
        import pptx  # noqa: F401
    except ImportError:
        print("python-pptx is required: pip install python-pptx")
        return 1

    with open(md_path, "r", encoding="utf-8-sig") as f:
        text = f.read()
    try:
        deck = parse_md(text)
    except MdContractError as e:
        print("md contract violation: %s" % e)
        return 1

    build_pptx(deck, out_path)
    problems = verify_pptx(deck, out_path)
    if problems:
        for p in problems:
            print("MISMATCH %s" % p)
        return 1
    print(
        "OK: %d slides (cover + %d), titles and notes match -> %s"
        % (1 + len(deck.slides), len(deck.slides), os.path.basename(out_path))
    )
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
