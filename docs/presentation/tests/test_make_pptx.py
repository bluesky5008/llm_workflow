# make_pptx.py 검증 테스트 — 외부 프레임워크 없이 단독 실행: python tests/test_make_pptx.py
# 견본 md(표지+본문 2장, 중첩 불릿·코드 블록·노트)로 파서·생성기·자동 대조를 검증한다.
import os
import sys
import tempfile

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except AttributeError:
    pass

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import make_pptx  # noqa: E402

SAMPLE = """# 견본 덱

- 부제 첫 줄
- 부제 둘째 줄

> 노트: 표지 노트입니다.

## 첫 슬라이드

- 항목 하나
  - 하위 항목

```
code line 1
code line 2
```

- 코드 뒤 항목

> 노트: 노트 한 줄.
> 둘째 줄.

## 둘째 슬라이드

- 단일 항목
"""


def test_parse():
    deck = make_pptx.parse_md(SAMPLE)
    assert deck.title == "견본 덱", deck.title
    assert deck.cover_lines == ["부제 첫 줄", "부제 둘째 줄"], deck.cover_lines
    assert deck.cover_notes == "표지 노트입니다.", repr(deck.cover_notes)
    assert len(deck.slides) == 2, len(deck.slides)

    s1 = deck.slides[0]
    assert s1.title == "첫 슬라이드", s1.title
    assert s1.blocks == [
        ("bullet", 0, "항목 하나"),
        ("bullet", 1, "하위 항목"),
        ("code", "code line 1\ncode line 2"),
        ("bullet", 0, "코드 뒤 항목"),
    ], s1.blocks
    assert s1.notes == "노트 한 줄.\n둘째 줄.", repr(s1.notes)

    s2 = deck.slides[1]
    assert s2.title == "둘째 슬라이드", s2.title
    assert s2.blocks == [("bullet", 0, "단일 항목")], s2.blocks
    assert s2.notes == "", repr(s2.notes)


def test_parse_rejects_unsupported():
    # 계약 밖 요소(표)는 줄 번호와 함께 실패해야 한다 (조용한 누락 금지).
    bad = "# 덱\n\n## 장\n\n| a | b |\n"
    try:
        make_pptx.parse_md(bad)
    except make_pptx.MdContractError as e:
        assert "5" in str(e), str(e)
    else:
        raise AssertionError("MdContractError not raised for table line")


def test_build_and_verify():
    deck = make_pptx.parse_md(SAMPLE)
    out = os.path.join(tempfile.mkdtemp(prefix="wf-pptx-test-"), "sample.pptx")
    make_pptx.build_pptx(deck, out)
    assert os.path.isfile(out), out

    problems = make_pptx.verify_pptx(deck, out)
    assert problems == [], problems

    from pptx import Presentation

    prs = Presentation(out)
    slides = list(prs.slides)
    assert len(slides) == 3, len(slides)
    assert slides[0].shapes.title.text == "견본 덱"
    assert slides[1].shapes.title.text == "첫 슬라이드"
    assert slides[2].shapes.title.text == "둘째 슬라이드"

    body = slides[1].placeholders[1].text_frame.text
    assert "항목 하나" in body, body
    assert "code line 1" in body, body
    assert slides[1].notes_slide.notes_text_frame.text == "노트 한 줄.\n둘째 줄."
    assert slides[0].notes_slide.notes_text_frame.text == "표지 노트입니다."


def test_verify_detects_mismatch():
    # 대조 함수가 실제 불일치를 잡아내는지 확인한다 (항상 빈 목록을 돌려주는 구현 방지).
    deck = make_pptx.parse_md(SAMPLE)
    out = os.path.join(tempfile.mkdtemp(prefix="wf-pptx-test-"), "sample.pptx")
    make_pptx.build_pptx(deck, out)
    deck.slides[0].title = "다른 제목"
    problems = make_pptx.verify_pptx(deck, out)
    assert problems, "verify_pptx must report a title mismatch"


def main():
    tests = [
        test_parse,
        test_parse_rejects_unsupported,
        test_build_and_verify,
        test_verify_detects_mismatch,
    ]
    failed = 0
    for t in tests:
        try:
            t()
            print("PASS %s" % t.__name__)
        except Exception as e:  # noqa: BLE001 — 단독 하니스: 모든 실패를 보고
            failed += 1
            print("FAIL %s: %r" % (t.__name__, e))
    print("%d/%d passed" % (len(tests) - failed, len(tests)))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
